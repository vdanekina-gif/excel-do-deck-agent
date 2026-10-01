"""
excel-to-deck agent — reads client_performance_data.xlsx and assembles a
management-review .pptx: one portfolio summary slide + one slide per client.

Usage: python3 scripts/build_deck.py
Reads:  client_performance_data.xlsx  (in the project root)
Writes: Meridian_Advisory_Portfolio_Review.pptx (in the project root)
"""
import datetime
from pathlib import Path

import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
INPUT_XLSX = ROOT / "client_performance_data.xlsx"
OUTPUT_PPTX = ROOT / "Meridian_Advisory_Portfolio_Review.pptx"

COMPANY_NAME = "Meridian Advisory Group"
REPORT_LABEL = "Portfolio Review — August 2026"
REPORT_DATE = datetime.date(2026, 8, 31)
RENEWAL_RISK_DAYS = 183          # ~6 months
CSAT_OVERDUE_DAYS = 183           # ~6 months
HIGH_ATTRITION = 0.15

# ---- palette (Midnight Executive) -------------------------------------
NAVY = RGBColor(0x1E, 0x27, 0x61)
ICE = RGBColor(0xE8, 0xEE, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5A, 0x64, 0x72)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xB5, 0x6A, 0x0B)
RED = RGBColor(0xC6, 0x28, 0x28)
GRAY = RGBColor(0x8C, 0x8C, 0x8C)

FONT = "Calibri"
HEADER_FONT = "Cambria"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_W = SLIDE_W - 2 * MARGIN


# ------------------------------------------------------------------ data --
def load_clients(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Clients"]
    headers = [c.value for c in ws[1]]
    clients = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[1] is None:
            break  # stop at the notes row (merged, only column A populated)
        d = dict(zip(headers, row))
        clients.append(d)
    return clients


def days_between(d1, d2):
    if isinstance(d1, datetime.datetime):
        d1 = d1.date()
    if isinstance(d2, datetime.datetime):
        d2 = d2.date()
    return (d1 - d2).days


def compute_portfolio(clients):
    rev_ty = sum(c["Revenue this year (projected, $)"] for c in clients)
    rev_ly = sum(c["Revenue last year ($)"] for c in clients)
    gm_ty = sum(c["Revenue this year (projected, $)"] * c["GM this year"] for c in clients) / rev_ty
    gm_ly = sum(c["Revenue last year ($)"] * c["GM last year"] for c in clients) / rev_ly
    headcount = sum(c["Headcount"] for c in clients)
    non_billable = sum(c["Non-billable resources"] for c in clients)
    util_vals = [c["Utilization %"] for c in clients if c["Utilization %"] is not None]
    projects = sum(c["Number of projects"] for c in clients)
    project_based = [c for c in clients if c["Engagement type"] == "Project-based"]
    on_track = sum(1 for c in project_based if c["Delivery status"] == "On track")
    at_risk = sum(1 for c in project_based if c["Delivery status"] == "At risk")
    delayed = sum(1 for c in project_based if c["Delivery status"] == "Delayed")
    open_risks = sum(c["Open risks"] for c in clients)
    nps_avg = sum(c["NPS"] for c in clients) / len(clients)
    enps_avg = sum(c["eNPS"] for c in clients) / len(clients)
    attr_avg = sum(c["Attrition %"] for c in clients) / len(clients)
    return dict(
        rev_ty=rev_ty, rev_ly=rev_ly, yoy=rev_ty / rev_ly - 1,
        gm_ty=gm_ty, gm_ly=gm_ly, gm_delta=gm_ty - gm_ly,
        headcount=headcount, non_billable=non_billable,
        util_avg=(sum(util_vals) / len(util_vals)) if util_vals else None,
        util_n=len(util_vals), staff_aug_n=sum(1 for c in clients if c["Engagement type"] == "Staff augmentation"),
        projects=projects, on_track=on_track, at_risk=at_risk, delayed=delayed,
        project_based_n=len(project_based),
        open_risks=open_risks, nps_avg=nps_avg, enps_avg=enps_avg, attr_avg=attr_avg,
    )


# --------------------------------------------------------------- helpers --
def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def add_textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return box, tf


def set_run(run, text, size, color, bold=False, font=FONT, italic=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def new_para(tf, first=False):
    if first:
        return tf.paragraphs[0]
    return tf.add_paragraph()


def money(v):
    return f"${v/1_000_000:,.2f}M" if abs(v) >= 1_000_000 else f"${v:,.0f}"


def pct(v, digits=1):
    return f"{v*100:.{digits}f}%"


def delta_pct(v, digits=1):
    sign = "+" if v >= 0 else ""
    return f"{sign}{v*100:.{digits}f}%"


def delta_pts(v, digits=1):
    sign = "+" if v >= 0 else ""
    return f"{sign}{v*100:.{digits}f} pts"


def trend_color(v):
    return GREEN if v >= 0 else RED


def add_background(slide, color=WHITE):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    set_fill(bg, color)
    bg.shadow.inherit = False
    # send to back
    spTree = slide.shapes._spTree
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)
    return bg


def add_footer(slide, page_no, total_pages):
    box, tf = add_textbox(slide, MARGIN, SLIDE_H - Inches(0.45), CONTENT_W, Inches(0.3))
    p = new_para(tf, first=True)
    r = p.add_run()
    set_run(r, f"{COMPANY_NAME}  ·  {REPORT_LABEL}  ·  synthetic demo data", 9, GRAY, italic=True)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    set_run(r2, f"{page_no} / {total_pages}", 9, GRAY)
    # merge into one line using tabs would be cleaner; keep simple two-line-safe box
    box.text_frame.word_wrap = False


def kpi_card(slide, left, top, width, height, label, value, sub=None, sub_color=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.adjustments[0] = 0.06
    set_fill(card, ICE)

    pad = Inches(0.18)
    box, tf = add_textbox(slide, left + pad, top + pad, width - 2 * pad, height - 2 * pad)
    p = new_para(tf, first=True)
    r = p.add_run()
    set_run(r, label.upper(), 10.5, MUTED, bold=True, font=FONT)

    p2 = tf.add_paragraph()
    p2.space_before = Pt(4)
    r2 = p2.add_run()
    set_run(r2, value, 30, NAVY, bold=True, font=HEADER_FONT)

    if sub:
        p3 = tf.add_paragraph()
        p3.space_before = Pt(2)
        r3 = p3.add_run()
        set_run(r3, sub, 10.5, sub_color or MUTED, bold=False)


def chip(slide, left, top, width, height, label, value, value_color=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.adjustments[0] = 0.10
    set_fill(card, WHITE)
    card.line.color.rgb = ICE
    card.line.width = Pt(1.25)
    card.shadow.inherit = False

    pad = Inches(0.14)
    box, tf = add_textbox(slide, left + pad, top + pad, width - 2 * pad, height - 2 * pad)
    p = new_para(tf, first=True)
    r = p.add_run()
    set_run(r, label.upper(), 9, MUTED, bold=True)

    p2 = tf.add_paragraph()
    p2.space_before = Pt(3)
    r2 = p2.add_run()
    set_run(r2, value, 17, value_color or NAVY, bold=True, font=HEADER_FONT)


def badge(slide, left, top, text, color):
    w, h = Inches(1.55), Inches(0.32)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    shape.adjustments[0] = 0.5
    set_fill(shape, color)
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, text, 10.5, WHITE, bold=True)
    return shape


def section_card(slide, left, top, width, height, title):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.adjustments[0] = 0.04
    set_fill(card, ICE)

    pad = Inches(0.22)
    box, tf = add_textbox(slide, left + pad, top + pad, width - 2 * pad, Inches(0.3))
    p = new_para(tf, first=True)
    r = p.add_run()
    set_run(r, title.upper(), 12, NAVY, bold=True, font=HEADER_FONT)
    return left + pad, top + pad + Inches(0.42), width - 2 * pad, height - pad - Inches(0.42) - pad


def stat_line(tf, first, label, value, value_color=NAVY, flag=None):
    p = new_para(tf, first=first)
    p.space_after = Pt(6)
    r1 = p.add_run()
    set_run(r1, f"{label}:  ", 11.5, MUTED)
    r2 = p.add_run()
    set_run(r2, value, 11.5, value_color, bold=True)
    if flag:
        r3 = p.add_run()
        set_run(r3, f"   {flag}", 10.5, AMBER, italic=True, bold=True)


# --------------------------------------------------------------- slide 1 --
def build_summary_slide(prs, clients, agg, page_no, total_pages):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)

    box, tf = add_textbox(slide, MARGIN, Inches(0.45), CONTENT_W, Inches(0.55))
    p = new_para(tf, first=True)
    r = p.add_run()
    set_run(r, COMPANY_NAME, 30, NAVY, bold=True, font=HEADER_FONT)

    box2, tf2 = add_textbox(slide, MARGIN, Inches(1.02), CONTENT_W, Inches(0.4))
    p = new_para(tf2, first=True)
    r = p.add_run()
    set_run(r, f"{REPORT_LABEL}  ·  4 enterprise accounts  ·  portfolio-wide snapshot", 13, MUTED)

    top1 = Inches(1.75)
    h1 = Inches(1.7)
    n1 = 4
    gap1 = Inches(0.3)
    w1 = Emu(int((CONTENT_W - gap1 * (n1 - 1)) / n1))

    cards = [
        ("Total revenue (projected)", money(agg["rev_ty"]),
         f"vs {money(agg['rev_ly'])} LY  ·  {delta_pct(agg['yoy'])}", trend_color(agg["yoy"])),
        ("Blended gross margin", pct(agg["gm_ty"]),
         f"vs {pct(agg['gm_ly'])} LY  ·  {delta_pts(agg['gm_delta'])}", trend_color(agg["gm_delta"])),
        ("Total headcount", str(agg["headcount"]),
         f"{agg['non_billable']} non-billable", MUTED),
        (f"Avg utilization ({agg['staff_aug_n']} staff-aug accts)",
         pct(agg["util_avg"]) if agg["util_avg"] is not None else "N/A",
         "per-head billing model only", MUTED),
    ]
    for i, (label, value, sub, sub_color) in enumerate(cards):
        left = MARGIN + i * (w1 + gap1)
        kpi_card(slide, left, top1, w1, h1, label, value, sub, sub_color)

    top2 = top1 + h1 + Inches(0.35)
    h2 = Inches(1.35)
    n2 = 5
    gap2 = Inches(0.25)
    w2 = Emu(int((CONTENT_W - gap2 * (n2 - 1)) / n2))

    delivery_txt = f"{agg['on_track']} on-track / {agg['at_risk']} at-risk"
    if agg["delayed"]:
        delivery_txt += f" / {agg['delayed']} delayed"
    chips = [
        ("Delivery health (proj.-based)", delivery_txt, AMBER if agg["at_risk"] or agg["delayed"] else GREEN),
        ("Open risks (portfolio)", str(agg["open_risks"]), AMBER if agg["open_risks"] >= 8 else NAVY),
        ("Avg NPS", f"{agg['nps_avg']:.1f}", NAVY),
        ("Avg eNPS", f"{agg['enps_avg']:.1f}", NAVY),
        ("Avg attrition", pct(agg["attr_avg"]), AMBER if agg["attr_avg"] >= HIGH_ATTRITION else NAVY),
    ]
    for i, (label, value, color) in enumerate(chips):
        left = MARGIN + i * (w2 + gap2)
        chip(slide, left, top2, w2, h2, label, value, color)

    top3 = top2 + h2 + Inches(0.3)
    box3, tf3 = add_textbox(slide, MARGIN, top3, CONTENT_W, Inches(0.9))
    p = new_para(tf3, first=True)
    r = p.add_run()
    flagged = []
    for c in clients:
        days_to_end = days_between(c["Contract end"], REPORT_DATE)
        if 0 <= days_to_end <= RENEWAL_RISK_DAYS:
            flagged.append(f"{c['Client']} (renewal due)")
        if c["Attrition %"] >= HIGH_ATTRITION:
            flagged.append(f"{c['Client']} (high attrition)")
    if flagged:
        set_run(r, "Flagged for attention this month:  ", 11.5, MUTED, bold=True)
        r2 = p.add_run()
        set_run(r2, "  ·  ".join(sorted(set(flagged))), 11.5, AMBER, bold=True)
    else:
        set_run(r, "No accounts flagged for attention this month.", 11.5, MUTED, italic=True)

    add_footer(slide, page_no, total_pages)
    return slide


# ---------------------------------------------------------- client slide --
def build_client_slide(prs, client, page_no, total_pages):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)

    box, tf = add_textbox(slide, MARGIN, Inches(0.45), CONTENT_W, Inches(0.5))
    p = new_para(tf, first=True)
    r = p.add_run()
    set_run(r, client["Client"], 27, NAVY, bold=True, font=HEADER_FONT)

    since = client["Since"]
    since_txt = since.strftime("%b %Y") if isinstance(since, (datetime.date, datetime.datetime)) else str(since)
    end = client["Contract end"]
    end_txt = end.strftime("%b %Y") if isinstance(end, (datetime.date, datetime.datetime)) else str(end)
    days_to_end = days_between(end, REPORT_DATE)
    renewal_flag = " (renewal due within 6 months)" if 0 <= days_to_end <= RENEWAL_RISK_DAYS else ""

    box2, tf2 = add_textbox(slide, MARGIN, Inches(0.98), CONTENT_W, Inches(0.4))
    p = new_para(tf2, first=True)
    r = p.add_run()
    set_run(r, f"{client['Engagement type']}  ·  client since {since_txt}  ·  contract end {end_txt}", 12.5, MUTED)
    if renewal_flag:
        r2 = p.add_run()
        set_run(r2, renewal_flag, 12.5, AMBER, bold=True, italic=True)

    top = Inches(1.7)
    bottom_margin = Inches(0.55)
    avail_h = SLIDE_H - top - bottom_margin
    gap = Inches(0.35)
    card_w = Emu(int((CONTENT_W - gap) / 2))
    card_h = Emu(int((avail_h - gap) / 2))

    positions = {
        "financial": (MARGIN, top),
        "resourcing": (MARGIN + card_w + gap, top),
        "delivery": (MARGIN, top + card_h + gap),
        "relationship": (MARGIN + card_w + gap, top + card_h + gap),
    }

    # --- financial ---
    left, ctop = positions["financial"]
    cl, ct, cw, ch = section_card(slide, left, ctop, card_w, card_h, "Financial")
    box, tf = add_textbox(slide, cl, ct, cw, ch)
    yoy = client["Revenue this year (projected, $)"] / client["Revenue last year ($)"] - 1
    gm_delta = client["GM this year"] - client["GM last year"]
    stat_line(tf, True, "Revenue (projected)", money(client["Revenue this year (projected, $)"]))
    stat_line(tf, False, "vs last year", f"{money(client['Revenue last year ($)'])}  ({delta_pct(yoy)})",
               value_color=trend_color(yoy))
    stat_line(tf, False, "Gross margin", pct(client["GM this year"]))
    stat_line(tf, False, "vs last year", f"{pct(client['GM last year'])}  ({delta_pts(gm_delta)})",
               value_color=trend_color(gm_delta))

    # --- resourcing ---
    left, ctop = positions["resourcing"]
    cl, ct, cw, ch = section_card(slide, left, ctop, card_w, card_h, "Resourcing")
    box, tf = add_textbox(slide, cl, ct, cw, ch)
    util = client["Utilization %"]
    util_txt = pct(util) if util is not None else "N/A — client-managed"
    stat_line(tf, True, "Headcount", str(client["Headcount"]))
    stat_line(tf, False, "Non-billable resources", str(client["Non-billable resources"]))
    stat_line(tf, False, "Utilization", util_txt, value_color=NAVY if util is not None else GRAY)
    stat_line(tf, False, "Active projects", str(client["Number of projects"]))

    # --- delivery & risk ---
    left, ctop = positions["delivery"]
    cl, ct, cw, ch = section_card(slide, left, ctop, card_w, card_h, "Delivery & Risk")
    box, tf = add_textbox(slide, cl, ct, cw, Inches(0.35))
    status = client["Delivery status"]
    stat_line(tf, True, "Open risks", str(client["Open risks"]),
               value_color=AMBER if client["Open risks"] >= 3 else NAVY)
    status_color_map = {"On track": GREEN, "At risk": AMBER, "Delayed": RED}
    badge_color = status_color_map.get(status, GRAY)
    badge_text = status if status in status_color_map else "N/A"

    box_lbl, tf_lbl = add_textbox(slide, cl, ct + Inches(0.38), cw, Inches(0.25))
    p_label = new_para(tf_lbl, first=True)
    r = p_label.add_run()
    set_run(r, "Delivery status:", 11.5, MUTED)
    badge(slide, cl, ct + Inches(0.62), badge_text, badge_color)
    if status not in status_color_map:
        box_note, tf_note = add_textbox(slide, cl, ct + Inches(0.98), cw, Inches(0.55))
        p = new_para(tf_note, first=True)
        r = p.add_run()
        set_run(r, "Client owns delivery under this engagement model.", 10, GRAY, italic=True)

    # --- relationship ---
    left, ctop = positions["relationship"]
    cl, ct, cw, ch = section_card(slide, left, ctop, card_w, card_h, "Relationship Health")
    box, tf = add_textbox(slide, cl, ct, cw, ch)
    csat_date = client["Last NPS date"]
    csat_days = days_between(REPORT_DATE, csat_date)
    csat_flag = "⚠ overdue" if csat_days > CSAT_OVERDUE_DAYS else None
    csat_txt = csat_date.strftime("%b %Y") if isinstance(csat_date, (datetime.date, datetime.datetime)) else str(csat_date)
    attr = client["Attrition %"]
    stat_line(tf, True, "NPS", f"{client['NPS']}  (as of {csat_txt})",
               value_color=AMBER if csat_flag else NAVY, flag=csat_flag)
    stat_line(tf, False, "eNPS (account team)", str(client["eNPS"]))
    stat_line(tf, False, "Attrition (account team)", pct(attr),
               value_color=AMBER if attr >= HIGH_ATTRITION else NAVY,
               flag="⚠ high" if attr >= HIGH_ATTRITION else None)

    add_footer(slide, page_no, total_pages)
    return slide


# --------------------------------------------------------------------- main
def main():
    clients = load_clients(INPUT_XLSX)
    agg = compute_portfolio(clients)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    total_pages = 1 + len(clients)
    build_summary_slide(prs, clients, agg, 1, total_pages)
    for i, client in enumerate(clients, start=2):
        build_client_slide(prs, client, i, total_pages)

    prs.save(OUTPUT_PPTX)
    print(f"saved {OUTPUT_PPTX}")


if __name__ == "__main__":
    main()
